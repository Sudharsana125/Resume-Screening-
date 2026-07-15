"""
Resume Parsing Module
======================
Extracts raw text from PDF and DOCX resumes, plus lightweight contact-info
extraction (email, phone) via regex. This module is intentionally decoupled
from NLP/skill-extraction so it can be tested or swapped independently.
"""
import re
import os
from dataclasses import dataclass, field
from typing import Optional

import pdfplumber
import docx  # python-docx


EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")

# Matches common phone formats: +91 98765 43210, (123) 456-7890, 123-456-7890 etc.
# Built as a run of digit-groups separated by space/dot/dash, with an optional
# leading country code and optional parenthesized area code.
PHONE_RE = re.compile(
    r"(?:\+\d{1,3}[\s.-]?)?"
    r"(?:\(\d{2,4}\)[\s.-]?)?"
    r"\d{3,5}[\s.-]?\d{3,5}(?:[\s.-]?\d{2,4})?"
)

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class UnsupportedFileTypeError(Exception):
    pass


class EmptyResumeError(Exception):
    """Raised when no extractable text is found (e.g. scanned image PDF)."""
    pass


@dataclass
class ParsedResume:
    file_name: str
    file_type: str
    raw_text: str
    email: Optional[str] = None
    phone: Optional[str] = None
    sections: dict = field(default_factory=dict)


def _extract_text_from_pdf(path: str) -> str:
    text_chunks = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text() or ""
            text_chunks.append(page_text)
    return "\n".join(text_chunks).strip()


def _extract_text_from_docx(path: str) -> str:
    document = docx.Document(path)
    paragraphs = [p.text for p in document.paragraphs]

    # Also pull text out of tables, since many resumes use table layouts
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)

    return "\n".join(p for p in paragraphs if p.strip()).strip()


def _extract_contact_info(text: str) -> tuple[Optional[str], Optional[str]]:
    email_match = EMAIL_RE.search(text)
    email = email_match.group(0) if email_match else None

    phone = None
    for match in PHONE_RE.finditer(text):
        candidate = match.group(0)
        digit_count = sum(c.isdigit() for c in candidate)
        # avoid false positives like page numbers / years by requiring
        # a realistic phone-number digit count
        if 7 <= digit_count <= 15:
            phone = candidate.strip()
            break

    return email, phone


# Common resume section headers used for naive section splitting.
_SECTION_HEADERS = [
    "experience", "work experience", "professional experience",
    "education", "skills", "technical skills", "projects",
    "certifications", "summary", "objective", "achievements",
]


def _split_sections(text: str) -> dict:
    """
    Very lightweight section splitter: scans line-by-line and treats a line
    matching a known header (short line, no trailing punctuation) as the
    start of a new section. Good enough for downstream heuristics like
    education/experience extraction; NOT meant to be a perfect parser.
    """
    sections: dict[str, list[str]] = {}
    current = "header"
    sections[current] = []

    for line in text.splitlines():
        stripped = line.strip()
        lower = stripped.lower().strip(":")
        if lower in _SECTION_HEADERS and len(stripped) < 40:
            current = lower
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line)

    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


def parse_resume(file_path: str) -> ParsedResume:
    """
    Main entry point. Detects file type by extension and dispatches to the
    correct extractor.

    Raises:
        UnsupportedFileTypeError: if extension isn't .pdf or .docx
        EmptyResumeError: if no text could be extracted (e.g. scanned image)
    """
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"Unsupported file type '{ext}'. Supported types: {SUPPORTED_EXTENSIONS}"
        )

    if ext == ".pdf":
        raw_text = _extract_text_from_pdf(file_path)
        file_type = "pdf"
    else:
        raw_text = _extract_text_from_docx(file_path)
        file_type = "docx"

    if not raw_text.strip():
        raise EmptyResumeError(
            "No text could be extracted from this resume. It may be a "
            "scanned image PDF that requires OCR."
        )

    email, phone = _extract_contact_info(raw_text)
    sections = _split_sections(raw_text)

    return ParsedResume(
        file_name=os.path.basename(file_path),
        file_type=file_type,
        raw_text=raw_text,
        email=email,
        phone=phone,
        sections=sections,
    )
